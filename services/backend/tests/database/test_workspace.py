"""Real runtime-role tests; requests do not run as the migration owner."""

import io
import json
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl
import pyotp
import pytest
import sqlalchemy as sa
from fastapi.testclient import TestClient
from pydantic import SecretStr
from sqlalchemy.engine import Engine, make_url

from app.common import digest
from app.config import Environment, Settings
from app.database.runtime import Database
from app.domain.models import ProcessingJob
from app.main import create_app
from app.worker import process_job

pytestmark = pytest.mark.database


@pytest.fixture
def workspace_settings(database_engine: Engine, tmp_path: Path) -> Settings:
    local = Settings()
    assert local.runtime_database_url and local.auth_database_url
    return local.model_copy(
        update={
            "app_env": Environment.TEST,
            "redis_url": None,
            "runtime_database_url": SecretStr(
                make_url(local.runtime_database_url.get_secret_value())
                .set(database=database_engine.url.database)
                .render_as_string(hide_password=False)
            ),
            "auth_database_url": SecretStr(
                make_url(local.auth_database_url.get_secret_value())
                .set(database=database_engine.url.database)
                .render_as_string(hide_password=False)
            ),
            "storage_backend": "filesystem",
            "storage_path": tmp_path / "objects",
            "extraction_provider": "local",
        }
    )


@pytest.fixture
def client(workspace_settings: Settings) -> Iterator[TestClient]:
    with TestClient(create_app(workspace_settings)) as client:
        client.headers["Origin"] = "http://testserver"
        yield client


def post(client: TestClient, path: str, body: Any) -> Any:
    headers = {
        "X-CSRF-Token": client.cookies.get("cgl_csrf", "") or "",
        "Idempotency-Key": str(__import__("uuid").uuid4()),
    }
    return client.post("/api/v1" + path, json=body, headers=headers)


def demo(client: TestClient) -> tuple[str, dict[str, Any]]:
    response = post(client, "/auth/demo", {})
    assert response.status_code == 201, response.text
    return response.json()["submission_id"], client.get("/api/v1/me").json()


def process_pending(configuration: Settings, actor: dict[str, Any]) -> None:
    db = Database(configuration)
    tenant, user = UUID(actor["tenant_id"]), UUID(actor["user_id"])
    with db.tenant(tenant, user) as session:
        ids = list(
            session.scalars(sa.select(ProcessingJob.id).where(ProcessingJob.status == "queued"))
        )
    db.close()
    for job_id in ids:
        process_job(configuration, tenant, user, job_id)


def reviewed_workspace(client: TestClient, configuration: Settings) -> tuple[str, dict[str, Any]]:
    submission_id, actor = demo(client)
    process_pending(configuration, actor)
    data = client.get("/api/v1/submissions/" + submission_id).json()
    assert all(job["status"] == "completed" for job in data["jobs"]), data["jobs"]
    assert any(group["conflicting"] for group in data["readiness"]["groups"])
    decisions = [
        {
            "group_key": group["key"],
            "fingerprint": group["fingerprint"],
            "action": "approve",
            "observation_id": group["observations"][0]["id"],
            "reason": "Synthetic test: broker verified this value against the source and confirmed the sales estimate.",
        }
        for group in data["readiness"]["groups"]
    ]
    response = post(
        client,
        f"/submissions/{submission_id}/review-decisions",
        {"expected_version": data["submission"]["row_version"], "decisions": decisions},
    )
    assert response.status_code == 200, response.text
    return submission_id, actor


def test_end_to_end_review_approval_export_and_staleness(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = reviewed_workspace(client, workspace_settings)
    data = client.get("/api/v1/submissions/" + submission_id).json()
    assert data["readiness"]["ready"], data["readiness"]["blockers"]
    generated = post(
        client,
        f"/submissions/{submission_id}/packages",
        {"expected_version": data["submission"]["row_version"]},
    )
    assert generated.status_code == 202, generated.text
    package_id = generated.json()["id"]
    process_pending(workspace_settings, actor)
    detail = client.get("/api/v1/packages/" + package_id).json()
    assert detail["package"]["status"] == "ready", detail
    previews = {}
    for format in ("pdf", "xlsx", "zip"):
        assert client.get(f"/api/v1/packages/{package_id}/artifacts/{format}").status_code == 409
        preview = client.get(f"/api/v1/packages/{package_id}/artifacts/{format}?preview=true")
        assert preview.status_code == 200
        previews[format] = preview.content
    assert previews["pdf"].startswith(b"%PDF")
    approved = post(
        client,
        f"/packages/{package_id}/approvals",
        {
            "manifest_hash": detail["package"]["manifest_hash"],
            "attestation": "I reviewed these exact synthetic files and approve this package for test export.",
        },
    )
    assert approved.status_code == 200, approved.text
    for format, content in previews.items():
        exported = client.get(f"/api/v1/packages/{package_id}/artifacts/{format}")
        assert exported.status_code == 200 and exported.content == content
    workbook = openpyxl.load_workbook(io.BytesIO(previews["xlsx"]))
    assert {"Exposure schedule", "Coverage schedule", "Loss schedule", "Source references"} <= set(
        workbook.sheetnames
    )
    assert workbook["Exposure schedule"].max_row > 1
    assert workbook["Coverage schedule"].max_row > 1
    assert workbook["Loss schedule"].max_row > 1
    assert workbook["Source references"].max_row > 1
    workbook.close()
    with zipfile.ZipFile(io.BytesIO(previews["zip"])) as archive:
        assert archive.read("submission.pdf") == previews["pdf"]
        assert archive.read("submission.xlsx") == previews["xlsx"]
        snapshot = json.loads(archive.read("reviewed-snapshot.json"))
        manifest = json.loads(archive.read("manifest.json"))
        assert digest(archive.read("reviewed-snapshot.json")) == manifest["revision_hash"]
        assert snapshot["facts"] and all(fact["decision_id"] for fact in snapshot["facts"])
        source_ids = {source["id"] for source in manifest["source_index"]}
        assert source_ids and all(
            set(fact["source_ids"]) <= source_ids for fact in snapshot["facts"]
        )
        assert len(manifest["documents"]) == 4
        assert len(manifest["files"]) == 6  # PDF, XLSX, and all four supporting documents.
        for file in manifest["files"]:
            assert digest(archive.read(file["filename"])) == file["sha256"]
    # Excluding a document is material and invalidates the approval for current export.
    doc_id = data["documents"][0]["id"]
    assert (
        post(
            client,
            f"/documents/{doc_id}/exclude",
            {"reason": "Replaced for a synthetic integrity test"},
        ).status_code
        == 200
    )
    assert client.get(f"/api/v1/packages/{package_id}/artifacts/pdf").status_code == 409
    assert (
        post(
            client,
            f"/packages/{package_id}/approvals",
            {
                "manifest_hash": detail["package"]["manifest_hash"],
                "attestation": "An old manifest cannot authorize a changed submission.",
            },
        ).status_code
        == 409
    )


def test_cross_tenant_records_sources_downloads_and_jobs_are_inaccessible(
    client: TestClient, workspace_settings: Settings
) -> None:
    first_id, actor = demo(client)
    process_pending(workspace_settings, actor)
    first = client.get("/api/v1/submissions/" + first_id).json()
    demo(client)  # New session belongs exclusively to another isolated fictional agency.
    for path in [
        f"/submissions/{first_id}",
        f"/insureds/{first['insured']['id']}",
        f"/documents/{first['documents'][0]['id']}/content",
        f"/sources/{first['sources'][0]['id']}",
        f"/submissions/{first_id}/audit",
    ]:
        assert client.get("/api/v1" + path).status_code == 404, path
    assert post(client, f"/jobs/{first['jobs'][0]['id']}/retry", {}).status_code == 404
    # Tenant headers cannot override the server-selected agency.
    response = client.get(
        "/api/v1/submissions/" + first_id, headers={"X-Tenant-ID": actor["tenant_id"]}
    )
    assert response.status_code == 404


def test_rls_fails_closed_and_transaction_scope_does_not_leak(
    client: TestClient, workspace_settings: Settings
) -> None:
    _, actor = demo(client)
    db = Database(workspace_settings)
    try:
        with db.runtime_engine.connect() as connection:
            assert connection.scalar(sa.text("SELECT count(*) FROM app.submission")) == 0
        with db.tenant(UUID(actor["tenant_id"]), UUID(actor["user_id"])) as session:
            assert session.scalar(sa.text("SELECT count(*) FROM app.submission")) == 1
        with db.runtime_engine.connect() as connection:
            assert connection.scalar(sa.text("SELECT count(*) FROM app.submission")) == 0
        with db.tenant(UUID(actor["tenant_id"]), UUID(int=1)) as session:
            assert session.scalar(sa.text("SELECT count(*) FROM app.submission")) == 0
    finally:
        db.close()


def test_authenticator_enrollment_csrf_logout_and_replay(client: TestClient) -> None:
    response = post(
        client,
        "/auth/register",
        {
            "email": "test-" + __import__("uuid").uuid4().hex + "@example.invalid",
            "password": "A fictional password 2026!",
            "name": "Broker Example",
            "agency": "Fictional MFA Agency",
        },
    )
    assert response.status_code == 200, response.text
    assert client.get("/api/v1/dashboard").status_code == 403
    secret = response.json()["secret"]
    code = pyotp.TOTP(secret).now()
    enrolled = post(client, "/auth/mfa/confirm", {"code": code})
    assert enrolled.status_code == 200 and len(enrolled.json()["recovery_codes"]) == 10
    actor = client.get("/api/v1/me").json()
    assert actor["mfa_verified"] and actor["auth_method"] == "password"
    assert client.post("/api/v1/auth/logout", json={}).status_code == 403
    assert post(client, "/auth/logout", {}).status_code == 200
    assert client.get("/api/v1/me").status_code == 401


def test_material_conflict_and_stale_review_cannot_generate(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = demo(client)
    process_pending(workspace_settings, actor)
    data = client.get("/api/v1/submissions/" + submission_id).json()
    response = post(
        client,
        f"/submissions/{submission_id}/packages",
        {"expected_version": data["submission"]["row_version"]},
    )
    assert response.status_code == 409
    group = data["readiness"]["groups"][0]
    response = post(
        client,
        f"/submissions/{submission_id}/review-decisions",
        {
            "expected_version": data["submission"]["row_version"],
            "decisions": [
                {
                    "group_key": group["key"],
                    "fingerprint": "old-evidence",
                    "action": "approve",
                    "observation_id": group["observations"][0]["id"],
                    "reason": "This must be rejected because evidence changed",
                }
            ],
        },
    )
    assert response.status_code == 409 and response.json()["detail"]["code"] == "EVIDENCE_CHANGED"


def test_role_changes_apply_immediately_and_last_admin_is_protected(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = demo(client)
    members = client.get("/api/v1/agency").json()["members"]
    me = next(member for member in members if member["id"] == actor["membership_id"])
    headers = {"X-CSRF-Token": client.cookies.get("cgl_csrf", "") or ""}
    response = client.patch(
        "/api/v1/memberships/" + me["id"],
        json={"expected_version": me["row_version"], "status": "active", "roles": ["viewer"]},
        headers=headers,
    )
    assert response.status_code == 409 and response.json()["detail"]["code"] == "LAST_ADMINISTRATOR"
    response = client.patch(
        "/api/v1/memberships/" + me["id"],
        json={"expected_version": me["row_version"], "status": "active", "roles": ["tenant_admin"]},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert client.get("/api/v1/agency").status_code == 200
    assert client.get("/api/v1/submissions/" + submission_id).status_code == 403
    assert post(client, "/insureds", {"name": "Unauthorized fictional insured"}).status_code == 403
    assert client.get("/api/v1/me").json()["roles"] == ["tenant_admin"]


def test_recovery_codes_are_one_use_and_old_totp_cannot_replay(client: TestClient) -> None:
    from uuid import uuid4

    email = f"recovery-{uuid4().hex}@example.invalid"
    password = "Fictional recovery password 2026!"
    registered = post(
        client,
        "/auth/register",
        {
            "email": email,
            "password": password,
            "name": "Recovery Broker",
            "agency": "Recovery Test Agency",
        },
    )
    code = pyotp.TOTP(registered.json()["secret"]).now()
    confirmed = post(client, "/auth/mfa/confirm", {"code": code})
    assert confirmed.status_code == 200
    assert post(client, "/auth/logout", {}).status_code == 200
    assert (
        post(
            client, "/auth/login", {"email": email, "password": password, "code": code}
        ).status_code
        == 401
    )
    recovery_code = confirmed.json()["recovery_codes"][0]
    recovered = post(
        client,
        "/auth/recover",
        {
            "email": email,
            "recovery_code": recovery_code,
            "new_password": "New fictional password 2027!",
        },
    )
    assert recovered.status_code == 200 and recovered.json()["mfa_enrollment_required"]
    assert client.get("/api/v1/dashboard").status_code == 403
    assert (
        post(
            client,
            "/auth/recover",
            {
                "email": email,
                "recovery_code": recovery_code,
                "new_password": "Another fictional password!",
            },
        ).status_code
        == 401
    )


def test_normal_upload_cannot_claim_synthetic_scan_bypass(
    client: TestClient, workspace_settings: Settings
) -> None:
    from app.demo import pdf
    from app.domain.models import DocumentVersion

    submission_id, actor = demo(client)
    response = client.post(
        f"/api/v1/submissions/{submission_id}/documents",
        files={
            "file": (
                "fictional-upload.pdf",
                pdf(["Fictional uploaded document for quarantine verification"]),
                "application/pdf",
            )
        },
        headers={
            "X-CSRF-Token": client.cookies.get("cgl_csrf", "") or "",
            "Idempotency-Key": "normal-upload-test",
        },
    )
    assert response.status_code == 202, response.text
    assert response.json()["scan_status"] == "quarantined"
    assert client.get("/api/v1/documents/" + response.json()["id"] + "/content").status_code == 409
    db = Database(workspace_settings)
    with db.tenant(UUID(actor["tenant_id"]), UUID(actor["user_id"])) as session:
        document = session.get(DocumentVersion, UUID(response.json()["id"]))
        assert document is not None and document.fixture_signature is None
    db.close()


def test_rendered_manifest_is_database_sealed(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = reviewed_workspace(client, workspace_settings)
    data = client.get("/api/v1/submissions/" + submission_id).json()
    generated = post(
        client,
        f"/submissions/{submission_id}/packages",
        {"expected_version": data["submission"]["row_version"]},
    )
    assert generated.status_code == 202
    package_id = generated.json()["id"]
    process_pending(workspace_settings, actor)
    db = Database(workspace_settings)
    try:
        with (
            pytest.raises(sa.exc.DBAPIError) as caught,
            db.tenant(UUID(actor["tenant_id"]), UUID(actor["user_id"])) as session,
        ):
            session.execute(
                sa.text("UPDATE app.package SET manifest='{}'::jsonb WHERE id=:id"),
                {"id": UUID(package_id)},
            )
        assert getattr(caught.value.orig, "sqlstate", None) == "55000"
    finally:
        db.close()


def test_material_account_edit_invalidates_submission_input(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, _ = reviewed_workspace(client, workspace_settings)
    data = client.get("/api/v1/submissions/" + submission_id).json()
    insured = data["insured"]
    payload = {
        key: insured[key]
        for key in (
            "name",
            "dba",
            "entity_type",
            "mailing_address",
            "operations",
            "external_reference",
            "status",
        )
    }
    payload.update(name="Fictional Renamed Entity LLC", expected_version=insured["row_version"])
    changed = client.patch(
        "/api/v1/insureds/" + insured["id"],
        json=payload,
        headers={"X-CSRF-Token": client.cookies.get("cgl_csrf", "") or ""},
    )
    assert changed.status_code == 200, changed.text
    updated = client.get("/api/v1/submissions/" + submission_id).json()
    assert updated["submission"]["input_version"] == data["submission"]["input_version"] + 1
    assert any(
        item["key"] == "missing:account_name_mismatch" for item in updated["readiness"]["blockers"]
    )


def test_material_omission_requires_exception_authority_and_cannot_waive_identity(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = reviewed_workspace(client, workspace_settings)
    data = client.get(f"/api/v1/submissions/{submission_id}").json()
    group = next(g for g in data["readiness"]["groups"] if g["field"] == "coverage.endorsement")
    # A supported adverse assertion cannot be disguised as an incorrect extraction.
    body = {
        "expected_version": data["submission"]["row_version"],
        "decisions": [
            {
                "group_key": group["key"],
                "fingerprint": group["fingerprint"],
                "action": "reject",
                "reason": "Attempt to drop supported endorsement",
            }
        ],
    }
    assert post(client, f"/submissions/{submission_id}/review-decisions", body).status_code == 422
    body["decisions"][0]["action"] = "omit"
    assert post(client, f"/submissions/{submission_id}/review-decisions", body).status_code == 200
    pending = client.get(f"/api/v1/submissions/{submission_id}").json()
    omission = next(
        b
        for b in pending["readiness"]["blockers"]
        if b["key"] == "missing:omission:" + group["key"]
    )
    assert not pending["readiness"]["ready"] and omission["waivable"]
    member = next(
        m
        for m in client.get("/api/v1/agency").json()["members"]
        if m["id"] == actor["membership_id"]
    )
    assert (
        client.patch(
            "/api/v1/memberships/" + member["id"],
            json={
                "expected_version": member["row_version"],
                "status": "active",
                "roles": ["tenant_admin", "preparer"],
            },
            headers={"X-CSRF-Token": client.cookies.get("cgl_csrf", "") or ""},
        ).status_code
        == 200
    )
    acceptance = {
        "expected_version": pending["submission"]["row_version"],
        "key": omission["key"],
        "fingerprint": omission["fingerprint"],
        "reason": "Material limitation accepted for fictional test",
    }
    assert post(client, f"/submissions/{submission_id}/exceptions", acceptance).status_code == 403
    member = next(
        m
        for m in client.get("/api/v1/agency").json()["members"]
        if m["id"] == actor["membership_id"]
    )
    assert (
        client.patch(
            "/api/v1/memberships/" + member["id"],
            json={
                "expected_version": member["row_version"],
                "status": "active",
                "roles": ["tenant_admin", "preparer", "approver"],
            },
            headers={"X-CSRF-Token": client.cookies.get("cgl_csrf", "") or ""},
        ).status_code
        == 200
    )
    accepted = post(client, f"/submissions/{submission_id}/exceptions", acceptance)
    assert accepted.status_code == 200, accepted.text
    assert any(
        e["key"] == omission["key"] and e["accepted_by"] == actor["user_id"]
        for e in accepted.json()["revision"]["snapshot"]["exceptions"]
    )
    # Omission of mandatory identity remains nonwaivable even for an approver.
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    name = next(g for g in current["readiness"]["groups"] if g["field"] == "insured.legal_name")
    assert (
        post(
            client,
            f"/submissions/{submission_id}/review-decisions",
            {
                "expected_version": current["submission"]["row_version"],
                "decisions": [
                    {
                        "group_key": name["key"],
                        "fingerprint": name["fingerprint"],
                        "action": "omit",
                        "reason": "Mandatory identity cannot disappear",
                    }
                ],
            },
        ).status_code
        == 200
    )
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert any(
        b["key"] == "required:insured.legal_name" and not b["waivable"]
        for b in current["readiness"]["blockers"]
    )


def test_excluded_adverse_and_failed_sources_remain_in_frozen_inventory(
    client: TestClient, workspace_settings: Settings
) -> None:
    import io
    import json
    import zipfile
    from uuid import uuid4

    from app.domain.models import DocumentVersion

    submission_id, actor = reviewed_workspace(client, workspace_settings)
    # Add a failed document directly through the normal scoped store service.
    db = Database(workspace_settings)
    with db.tenant(UUID(actor["tenant_id"]), UUID(actor["user_id"])) as session:
        template = session.scalar(
            sa.select(DocumentVersion).where(DocumentVersion.submission_id == UUID(submission_id))
        )
        assert template is not None
        # A quarantined failed source must remain disclosed after exclusion.
        failed = DocumentVersion(
            id=uuid4(),
            tenant_id=UUID(actor["tenant_id"]),
            submission_id=UUID(submission_id),
            filename="Unreadable-loss-runs.pdf",
            media_type="application/pdf",
            sha256="a" * 64,
            byte_size=10,
            object_key="unused-excluded-test",
            uploaded_by=UUID(actor["user_id"]),
            processing_status="failed",
        )
        session.add(failed)
        session.flush()
        failed_id = str(failed.id)
    db.close()
    assert (
        post(
            client,
            f"/documents/{failed_id}/exclude",
            {"reason": "Unreadable extra loss report; confirm limitation"},
        ).status_code
        == 200
    )
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    blocker = next(
        b for b in current["readiness"]["blockers"] if b["key"] == "missing:excluded:" + failed_id
    )
    assert (
        post(
            client,
            f"/submissions/{submission_id}/exceptions",
            {
                "expected_version": current["submission"]["row_version"],
                "key": blocker["key"],
                "fingerprint": blocker["fingerprint"],
                "reason": "Approver reviewed missing additional report and accepted limitation",
            },
        ).status_code
        == 200
    )
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert current["readiness"]["ready"], current["readiness"]["blockers"]
    generated = post(
        client,
        f"/submissions/{submission_id}/packages",
        {"expected_version": current["submission"]["row_version"]},
    )
    assert generated.status_code == 202, generated.text
    process_pending(workspace_settings, actor)
    archive = client.get(f"/api/v1/packages/{generated.json()['id']}/artifacts/zip?preview=true")
    assert archive.status_code == 200, archive.text
    with zipfile.ZipFile(io.BytesIO(archive.content)) as package:
        snapshot = json.loads(package.read("reviewed-snapshot.json"))
        excluded = next(d for d in snapshot["document_inventory"] if d["id"] == failed_id)
        assert (
            excluded["excluded_by"] == actor["user_id"]
            and excluded["processing_status"] == "failed"
        )
        assert excluded["included_in_exhibits"] is False
        assert not any(failed_id in name for name in package.namelist())
        assert b"Unreadable-loss-runs.pdf" in package.read("manifest.json")


def test_correction_requires_explicit_support_and_preserves_original_history(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, _ = reviewed_workspace(client, workspace_settings)
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    group = next(
        g for g in current["readiness"]["groups"] if g["field"] == "coverage.each_occurrence"
    )
    selected = group["observations"][0]
    decision = {
        "group_key": group["key"],
        "fingerprint": group["fingerprint"],
        "action": "correct",
        "observation_id": selected["id"],
        "value": "3000000",
        "context": selected["context"],
        "reason": "Fictional client corrected the requested limit",
    }
    body = {"expected_version": current["submission"]["row_version"], "decisions": [decision]}
    assert post(client, f"/submissions/{submission_id}/review-decisions", body).status_code == 422
    decision.update(
        support_type="document",
        support_source_ids=[selected["source_id"], *selected["supporting_source_ids"]],
    )
    assert post(client, f"/submissions/{submission_id}/review-decisions", body).status_code == 422
    decision.update(
        support_type="broker_attestation", attestation_origin="Fictional client call on 2026-10-29"
    )
    result = post(client, f"/submissions/{submission_id}/review-decisions", body)
    assert result.status_code == 200, result.text
    fact = next(
        f for f in result.json()["revision"]["snapshot"]["facts"] if f["group_key"] == group["key"]
    )
    assert fact["evidence_type"] == "broker_attestation" and not fact["source_ids"]
    assert selected["source_id"] in fact["original_source_ids"]


def test_concurrent_replayed_create_commits_once(client: TestClient) -> None:
    from concurrent.futures import ThreadPoolExecutor

    demo(client)
    headers = {
        "X-CSRF-Token": client.cookies.get("cgl_csrf", "") or "",
        "Idempotency-Key": "concurrent-create-test",
    }

    def create() -> Any:
        return client.post(
            "/api/v1/insureds", json={"name": "Concurrent Fictional LLC"}, headers=headers
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        first, second = list(executor.map(lambda _: create(), range(2)))
    assert first.status_code == second.status_code == 201
    assert first.json()["id"] == second.json()["id"]
    assert (
        len(
            [
                i
                for i in client.get("/api/v1/insureds").json()["items"]
                if i["name"] == "Concurrent Fictional LLC"
            ]
        )
        == 1
    )


def test_renewal_candidates_do_not_inherit_verification_and_require_current_confirmation(
    client: TestClient, workspace_settings: Settings
) -> None:
    prior_id, _ = reviewed_workspace(client, workspace_settings)
    prior = client.get(f"/api/v1/submissions/{prior_id}").json()
    created = post(
        client,
        "/submissions",
        {
            "insured_id": prior["insured"]["id"],
            "title": "Fictional renewal",
            "kind": "renewal",
            "prior_submission_id": prior_id,
            "intended_start": "2028-01-01",
            "intended_end": "2029-01-01",
        },
    )
    assert created.status_code == 201, created.text
    renewal_id = created.json()["id"]
    current = client.get(f"/api/v1/submissions/{renewal_id}").json()
    groups = current["readiness"]["groups"]
    assert groups and all(g.get("candidate") and not g["verified"] for g in groups)
    assert not current["readiness"]["facts"] and not current["packages"]
    candidate = next(g for g in groups if g["field"] == "insured.legal_name")
    decision = {
        "group_key": candidate["key"],
        "fingerprint": candidate["fingerprint"],
        "action": "correct",
        "value": candidate["candidate"]["value"],
        "context": candidate["context"],
        "reason": "Confirmed for current renewal",
        "support_type": "broker_attestation",
        "attestation_origin": "Fictional insured email confirming renewal identity",
    }
    response = post(
        client,
        f"/submissions/{renewal_id}/review-decisions",
        {"expected_version": current["submission"]["row_version"], "decisions": [decision]},
    )
    assert response.status_code == 200, response.text
    fact = response.json()["revision"]["snapshot"]["facts"][0]
    assert (
        fact["renewal_lineage"]["prior_revision_id"] == candidate["candidate"]["prior_revision_id"]
    )
    assert fact["evidence_type"] == "broker_attestation"
    assert not client.get(f"/api/v1/submissions/{renewal_id}").json()["readiness"]["ready"]


def test_terminal_submission_and_concurrent_task_controls(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = reviewed_workspace(client, workspace_settings)
    task = post(
        client, f"/submissions/{submission_id}/tasks", {"title": "Confirm fictional operations"}
    ).json()
    headers = {"X-CSRF-Token": client.cookies.get("cgl_csrf", "") or ""}
    url = "/api/v1/tasks/" + task["id"]
    assert (
        client.patch(
            url, json={"expected_version": 1, "status": "completed"}, headers=headers
        ).status_code
        == 422
    )
    saved = client.patch(
        url,
        json={
            "expected_version": 1,
            "due_date": "2026-12-01",
            "owner_membership_id": actor["membership_id"],
        },
        headers=headers,
    )
    assert saved.status_code == 200 and saved.json()["row_version"] == 2
    assert (
        client.patch(
            url, json={"expected_version": 1, "status": "open"}, headers=headers
        ).status_code
        == 409
    )
    completed = client.patch(
        url,
        json={
            "expected_version": 2,
            "status": "completed",
            "resolution_reason": "Client confirmed fictional operations by email",
        },
        headers=headers,
    )
    assert completed.status_code == 200 and completed.json()["resolution_reason"]
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    ready = post(
        client,
        f"/submissions/{submission_id}/status",
        {
            "expected_version": current["submission"]["row_version"],
            "status": "ready_for_package",
            "reason": "All fictional readiness checks completed",
        },
    )
    assert ready.status_code == 200, ready.text
    cancelled = post(
        client,
        f"/submissions/{submission_id}/status",
        {
            "expected_version": ready.json()["row_version"],
            "status": "cancelled",
            "reason": "Fictional placement cancelled by client",
        },
    )
    assert cancelled.status_code == 200
    assert (
        post(
            client,
            f"/submissions/{submission_id}/status",
            {
                "expected_version": cancelled.json()["row_version"],
                "status": "in_review",
                "reason": "Terminal submissions cannot be reopened",
            },
        ).status_code
        == 409
    )
    assert (
        post(
            client,
            f"/submissions/{submission_id}/tasks",
            {"title": "No edits after terminal transition"},
        ).status_code
        == 409
    )


def test_parse_checkpoint_survives_provider_failure_and_retry_reuses_sources(
    client: TestClient, workspace_settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.documents import DocumentError, classify, local_extract

    submission_id, actor = demo(client)
    provider_configuration = workspace_settings.model_copy(
        update={
            "extraction_provider": "gateway",
            "extraction_endpoint": "https://unused.example.invalid",
        }
    )

    def fail_provider(*args: Any) -> Any:
        raise DocumentError("PROVIDER_EXTRACTION_FAILED")

    monkeypatch.setattr("app.worker.provider_extract", fail_provider)
    process_pending(provider_configuration, actor)
    failed = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert failed["sources"] and all(j["status"] == "failed" for j in failed["jobs"])
    original_ids = {s["id"] for s in failed["sources"]}
    assert all(d["processing_inventory"]["parse_complete"] for d in failed["documents"])
    monkeypatch.setattr(
        "app.worker.provider_extract",
        lambda blocks, configuration: local_extract(blocks, classify(blocks)),
    )
    for job in failed["jobs"]:
        assert post(client, f"/jobs/{job['id']}/retry", {}).status_code == 202
    process_pending(provider_configuration, actor)
    retried = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert retried["readiness"]["groups"]
    assert {s["id"] for s in retried["sources"]}.issubset(original_ids)
    assert all(d["processing_status"] == "completed" for d in retried["documents"])


def test_replacement_preserves_old_assertions_and_needs_authorized_disposition(
    client: TestClient, workspace_settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.demo import pdf

    submission_id, actor = reviewed_workspace(client, workspace_settings)
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    prior = next(d for d in current["documents"] if d["filename"] == "Expiring-policy.pdf")
    data = pdf(
        [
            "FICTIONAL POLICY DECLARATIONS",
            "Legal insured name: Lakeshore Design Services LLC",
            "Policy number: RENEWED-1",
            "Policy effective: 2026-01-01",
            "Policy expiration: 2027-01-01",
            "Currency: USD",
            "Each occurrence limit: 1000000",
            "General aggregate limit: 2000000",
        ]
    )
    uploaded = client.post(
        f"/api/v1/submissions/{submission_id}/documents",
        files={"file": ("Replacement-policy.pdf", data, "application/pdf")},
        data={"supersedes_id": prior["id"]},
        headers={
            "X-CSRF-Token": client.cookies.get("cgl_csrf", "") or "",
            "Idempotency-Key": "legitimate-replacement-test",
        },
    )
    assert uploaded.status_code == 202, uploaded.text
    monkeypatch.setattr("app.worker.scan", lambda *args: "mock-scan-of-fictional-fixture")
    process_pending(workspace_settings, actor)
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    old = next(d for d in current["readiness"]["document_inventory"] if d["id"] == prior["id"])
    assert old["replaced_by"] == [uploaded.json()["id"]] and old["excluded_by"] == actor["user_id"]
    assert any(
        a["field"] == "coverage.endorsement" and "DEMO-001" in a["value"]
        for a in old["source_assertions"]
    )
    pending = [g for g in current["readiness"]["groups"] if not g["verified"]]
    decisions = [
        {
            "group_key": g["key"],
            "fingerprint": g["fingerprint"],
            "action": "approve",
            "observation_id": g["observations"][0]["id"],
            "reason": "Reviewed fictional replacement and source context",
        }
        for g in pending
    ]
    assert (
        post(
            client,
            f"/submissions/{submission_id}/review-decisions",
            {"expected_version": current["submission"]["row_version"], "decisions": decisions},
        ).status_code
        == 200
    )
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    blocker = next(
        b for b in current["readiness"]["blockers"] if b["key"] == "missing:excluded:" + prior["id"]
    )
    accepted = post(
        client,
        f"/submissions/{submission_id}/exceptions",
        {
            "expected_version": current["submission"]["row_version"],
            "key": blocker["key"],
            "fingerprint": blocker["fingerprint"],
            "reason": "Approver confirmed replacement adequacy and disclosed lost endorsement wording",
        },
    )
    assert accepted.status_code == 200, accepted.text
    assert client.get(f"/api/v1/submissions/{submission_id}").json()["readiness"]["ready"]


def test_repeated_entities_survive_review_and_frozen_package(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, _ = reviewed_workspace(client, workspace_settings)
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    decisions = [
        {
            "group_key": f"manual:repeat-{index}",
            "field": field,
            "action": "correct",
            "value": value,
            "context": {},
            "reason": "Client confirmed distinct fictional entity or location",
            "support_type": "broker_attestation",
            "attestation_origin": "Fictional insured emailed distinct DBA and location schedule",
        }
        for index, (field, value) in enumerate(
            [
                ("insured.dba", "First DBA"),
                ("insured.dba", "Second DBA"),
                ("location.address", "200 Another Fictional Way"),
                ("location.address", "300 Third Fictional Way"),
            ]
        )
    ]
    result = post(
        client,
        f"/submissions/{submission_id}/review-decisions",
        {"expected_version": current["submission"]["row_version"], "decisions": decisions},
    )
    assert result.status_code == 200, result.text
    facts = result.json()["revision"]["snapshot"]["facts"]
    assert len([f for f in facts if f["field"] == "insured.dba"]) == 2
    assert len([f for f in facts if f["field"] == "location.address"]) == 3
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert current["readiness"]["ready"], current["readiness"]["blockers"]
    assert (
        post(
            client,
            f"/submissions/{submission_id}/packages",
            {"expected_version": current["submission"]["row_version"]},
        ).status_code
        == 202
    )


def test_dashboard_counts_all_submissions_and_includes_dynamic_blockers(
    client: TestClient, workspace_settings: Settings
) -> None:
    from datetime import date

    from app.domain.models import Submission

    submission_id, actor = demo(client)
    original = client.get(f"/api/v1/submissions/{submission_id}").json()["submission"]
    db = Database(workspace_settings)
    with db.tenant(UUID(actor["tenant_id"]), UUID(actor["user_id"])) as session:
        for index in range(101):
            session.add(
                Submission(
                    tenant_id=UUID(actor["tenant_id"]),
                    insured_id=UUID(original["insured_id"]),
                    title=f"Fictional dashboard submission {index}",
                    kind="new_business",
                    intended_start=date(2027, 1, 1),
                    intended_end=date(2028, 1, 1),
                    assigned_membership_id=UUID(actor["membership_id"]),
                )
            )
    db.close()
    dashboard = client.get("/api/v1/dashboard?owner_membership_id=" + actor["membership_id"]).json()
    assert dashboard["metrics"]["active_submissions"] == dashboard["submission_count"] == 102
    assert len(dashboard["submissions"]) == 100
    assert len(dashboard["blocking_queue"]) == 100
    assert dashboard["metrics"]["open_discrepancies"] > 100


def test_exclusion_refreshes_conflicts_without_hiding_completeness(
    client: TestClient, workspace_settings: Settings
) -> None:
    submission_id, actor = demo(client)
    process_pending(workspace_settings, actor)
    data = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert any(i["kind"] == "conflict" and i["status"] == "open" for i in data["issues"])
    schedule = next(d for d in data["documents"] if d["filename"] == "Sales-schedule.xlsx")
    assert (
        post(
            client,
            f"/documents/{schedule['id']}/exclude",
            {
                "reason": "Fictional schedule superseded; material sales difference needs disposition"
            },
        ).status_code
        == 200
    )
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert not any(i["kind"] == "conflict" and i["status"] == "open" for i in current["issues"])
    assert any(
        b["key"] == "missing:excluded:" + schedule["id"] for b in current["readiness"]["blockers"]
    )
    assert not current["readiness"]["ready"]


def test_concurrent_upload_retry_creates_one_document_and_job(client: TestClient) -> None:
    from concurrent.futures import ThreadPoolExecutor

    from app.demo import pdf

    submission_id, _ = demo(client)
    data = pdf(["Legal insured name: Retry Fixture LLC"])
    headers = {
        "X-CSRF-Token": client.cookies.get("cgl_csrf", "") or "",
        "Idempotency-Key": "concurrent-upload-retry",
    }

    def upload_again() -> Any:
        return client.post(
            f"/api/v1/submissions/{submission_id}/documents",
            files={"file": ("Retry-fixture.pdf", data, "application/pdf")},
            headers=headers,
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        first, second = list(executor.map(lambda _: upload_again(), range(2)))
    assert first.status_code == second.status_code == 202
    assert first.json()["id"] == second.json()["id"]
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    assert len([d for d in current["documents"] if d["filename"] == "Retry-fixture.pdf"]) == 1
    assert len([j for j in current["jobs"] if j["document_id"] == first.json()["id"]]) == 1


def test_partial_policy_stays_incomplete_even_when_other_documents_supply_coverage(
    client: TestClient, workspace_settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.demo import pdf

    submission_id, actor = reviewed_workspace(client, workspace_settings)
    uploaded = client.post(
        f"/api/v1/submissions/{submission_id}/documents",
        files={
            "file": (
                "Partial-policy.pdf",
                pdf(
                    [
                        "FICTIONAL POLICY DECLARATIONS",
                        "Legal insured name: Lakeshore Design Services LLC",
                    ]
                ),
                "application/pdf",
            )
        },
        headers={
            "X-CSRF-Token": client.cookies.get("cgl_csrf", "") or "",
            "Idempotency-Key": "partial-policy-inventory",
        },
    )
    assert uploaded.status_code == 202
    monkeypatch.setattr("app.worker.scan", lambda *args: "mock-fictional-scan")
    process_pending(workspace_settings, actor)
    current = client.get(f"/api/v1/submissions/{submission_id}").json()
    issues = [
        b
        for b in current["readiness"]["blockers"]
        if b["key"].startswith("missing:policy_inventory:" + uploaded.json()["id"])
    ]
    assert (
        issues
        and "Each occurrence limit" in issues[0]["message"]
        and "Policy effective date" in issues[0]["message"]
    )
    assert not current["readiness"]["ready"]
