"""Regression cases from review.md: context, availability, provenance and partial PDFs."""

import io
from typing import Any
from uuid import uuid4

import pytest
from pypdf import PdfReader, PdfWriter

from app.demo import pdf
from app.documents import Block, DocumentError, local_extract, parse_pdf, supports_fact
from app.domain.cgl import group_key, validate_fact
from app.workflow import approved_facts


def blocks(*lines: str, page: int = 1) -> list[Block]:
    return [
        Block(line, {"kind": "pdf", "page": page, "line": i + 1}) for i, line in enumerate(lines)
    ]


def test_policy_context_never_flows_backward_or_across_pages() -> None:
    source = blocks(
        "Policy number: OLD-2026",
        "Currency: USD",
        "Each occurrence limit: 1,000,000",
        "Policy number: NEW-2027",
        "Currency: CAD",
        "Each occurrence limit: 2,000,000",
    )
    facts = [f for f in local_extract(source, "policy") if f.field == "coverage.each_occurrence"]
    assert [f.context["policy_reference"] for f in facts] == ["OLD-2026", "NEW-2027"]
    assert [f.context["currency"] for f in facts] == ["USD", "CAD"]
    assert source[3].id not in facts[0].source_ids
    assert source[0].id not in facts[1].source_ids
    ambiguous = local_extract(
        source + blocks("Each occurrence limit: 3,000,000", page=2), "policy"
    )[-1]
    assert ambiguous.availability == "unreadable" and ambiguous.value is None


def test_multiple_exposure_periods_statuses_and_currencies_stay_separate() -> None:
    facts = local_extract(
        blocks(
            "Currency: CAD",
            "Exposure period: 2026-01-01 to 2026-12-31",
            "Exposure status: actual",
            "Payroll: 100,000",
            "Currency: EUR",
            "Exposure period: 2027-01-01 to 2027-12-31",
            "Exposure status: estimate",
            "Payroll: 200,000",
        ),
        "exposure_schedule",
    )
    assert len(facts) == 2
    assert [
        (
            f.value,
            f.context["period_start"],
            f.context["role"],
            f.context["currency"],
            f.context["unit"],
        )
        for f in facts
    ] == [
        ("100000", "2026-01-01", "actual", "CAD", "money"),
        ("200000", "2027-01-01", "estimate", "EUR", "money"),
    ]


def test_mixed_document_section_resets_context() -> None:
    facts = local_extract(
        blocks(
            "Policy number: OLD",
            "Currency: USD",
            "Each occurrence limit: 1000",
            "CGL APPLICATION",
            "Requested each occurrence limit: 2000",
        ),
        "policy",
    )
    assert facts[-1].availability == "unreadable"
    assert "policy_reference" not in facts[-1].context


@pytest.mark.parametrize(
    "raw,state",
    [
        ("unknown", "unknown"),
        ("n/a", "not_applicable"),
        ("not supplied", "not_supplied"),
        ("not a number", "unreadable"),
    ],
)
def test_material_field_availability_is_preserved(raw: str, state: str) -> None:
    facts = local_extract(
        blocks("Legal insured name: Test LLC", "Currency: CAD", "Each occurrence limit: " + raw),
        "policy",
    )
    assert len(facts) == 2
    assert facts[-1].value is None and facts[-1].availability == state
    assert facts[-1].source_ids


def test_dbas_and_locations_have_repeated_row_identity() -> None:
    assert group_key("insured.dba", {}, "First") != group_key("insured.dba", {}, "Second")
    assert group_key("location.address", {}, "1 First St") != group_key(
        "location.address", {}, "2 Second St"
    )
    assert group_key("location.address", {"location_id": "building-1"}, "old") == group_key(
        "location.address", {"location_id": "building-1"}, "new"
    )
    assert group_key("insured.dba", {"entity_id": "entity-1"}, "old") == group_key(
        "insured.dba", {"entity_id": "entity-1"}, "new"
    )


@pytest.mark.parametrize("currency", ["CAD", "EUR"])
def test_foreign_exposures_reject_contradictory_units(currency: str) -> None:
    context = {
        "basis": "payroll",
        "unit": "USD",
        "currency": currency,
        "period_start": "2027-01-01",
        "period_end": "2027-12-31",
    }
    with pytest.raises(ValueError, match="unit"):
        validate_fact("exposure.amount", "1000", context)
    assert validate_fact("exposure.amount", "1000", {**context, "unit": "money"})[0] == "1000"


def test_ambiguous_dollar_sign_does_not_establish_currency() -> None:
    fact = local_extract(
        blocks("Exposure period: 2027-01-01 to 2027-12-31", "Payroll: $1000"), "exposure_schedule"
    )[0]
    assert fact.availability == "unreadable" and "currency" not in fact.context


def test_evidence_validator_requires_field_and_context_association() -> None:
    source = blocks("Policy number: P-1", "Currency: USD", "General aggregate limit: 2000000")
    context = {"role": "existing", "policy_reference": "P-1", "currency": "USD"}
    refs = [b.id for b in source]
    assert supports_fact(source, "coverage.general_aggregate", "2000000", context, refs)
    assert not supports_fact(source, "coverage.each_occurrence", "2000000", context, refs)
    assert not supports_fact(
        source, "coverage.general_aggregate", "2000000", {**context, "role": "requested"}, refs
    )
    assert not supports_fact(
        source,
        "coverage.general_aggregate",
        "2000000",
        {**context, "policy_reference": "INVENTED"},
        refs,
    )


def test_corrected_fact_has_separate_attestation_and_original_history() -> None:
    observation_id, source_id = str(uuid4()), str(uuid4())
    decision: dict[str, Any] = {
        "id": str(uuid4()),
        "field": "coverage.each_occurrence",
        "action": "correct",
        "value": "2000000",
        "context": {"currency": "CAD", "role": "requested"},
        "observation_id": observation_id,
        "reason": "Client corrected amount and requested context",
        "reviewer_id": str(uuid4()),
        "support": {
            "type": "broker_attestation",
            "origin": "Client call on 2026-10-29",
            "source_ids": [],
        },
    }
    facts = approved_facts(
        [
            {
                "key": "test",
                "verified": True,
                "decision": decision,
                "observations": [
                    {"id": observation_id, "source_id": source_id, "supporting_source_ids": []}
                ],
            }
        ]
    )
    assert facts[0]["source_ids"] == []
    assert facts[0]["original_source_ids"] == [source_id]
    assert facts[0]["evidence_type"] == "broker_attestation"
    assert facts[0]["attestation_origin"] == "Client call on 2026-10-29"


def test_pdf_retains_usable_pages_and_blank_page_inventory(monkeypatch: pytest.MonkeyPatch) -> None:
    writer = PdfWriter()
    writer.add_page(PdfReader(io.BytesIO(pdf(["Legal insured name: Usable LLC"]))).pages[0])
    writer.add_blank_page(width=612, height=792)
    writer.add_page(
        PdfReader(io.BytesIO(pdf(["Damaged page containing supplied material content"]))).pages[0]
    )
    stream = io.BytesIO()
    writer.write(stream)
    original = __import__("pypdf")._page.PageObject.extract_text

    def damaged(page: Any, *args: Any, **kwargs: Any) -> str:
        value = original(page, *args, **kwargs)
        return "" if "Damaged" in value else str(value)

    monkeypatch.setattr("pypdf._page.PageObject.extract_text", damaged)

    def failed_ocr(*args: Any) -> str:
        raise DocumentError("OCR_FAILED")

    monkeypatch.setattr("app.documents.ocr_page", failed_ocr)
    parsed, count = parse_pdf(stream.getvalue())
    assert count == 3
    assert local_extract(parsed, "application")[0].value == "Usable LLC"
    assert any(b.locator.get("page_status") == "blank" and b.locator["page"] == 2 for b in parsed)
    assert any(
        b.locator.get("page_status") == "unreadable" and b.locator["page"] == 3 for b in parsed
    )


def test_spreadsheet_cell_formats_merges_formulas_and_hidden_content_are_preserved() -> None:
    import openpyxl

    from app.documents import parse_xlsx, processing_inventory

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Currency", "CAD"])
    sheet.append(["Annual gross sales", "=1000*2"])
    sheet["B2"].number_format = "#,##0.00"
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "Merged source heading"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    parsed, _ = parse_xlsx(buffer.getvalue())
    row = next(b for b in parsed if b.locator.get("range") == "A2:B2")
    formula = row.locator["cells"][1]
    assert formula["formula"] == "=1000*2" and formula["cached_value"] is None
    assert formula["number_format"] == "#,##0.00" and formula["hidden_column"]
    assert row.locator["hidden_row"] and row.locator["warnings"]
    inventory = processing_inventory(parsed, local_extract(parsed, "exposure_schedule"))
    assert inventory["sheets"][0]["merged_ranges"] == ["A3:B3"]
    assert inventory["issues"]


def test_packages_disclose_excluded_documents_and_correction_support_in_all_formats() -> None:
    import json
    import zipfile

    import openpyxl

    from app.common import canonical_json, digest
    from app.rendering import render

    reviewer_id, source_id, document_id = str(uuid4()), str(uuid4()), str(uuid4())
    snapshot: dict[str, Any] = {
        "intended_start": "2027-01-01",
        "intended_end": "2028-01-01",
        "facts": [
            {
                "field": "exposure.amount",
                "value": "123456789.1234",
                "context": {
                    "basis": "payroll",
                    "unit": "money",
                    "currency": "CAD",
                    "role": "estimate",
                    "scope": "Long operation description " * 150,
                },
                "source_ids": [],
                "original_source_ids": [source_id],
                "evidence_type": "broker_attestation",
                "reviewer_id": reviewer_id,
                "attestation_origin": "Fictional client call on 2026-10-29",
                "decision_id": str(uuid4()),
                "reason": "Corrected amount based on client confirmation",
            }
        ],
        "exceptions": [
            {
                "key": "excluded-test",
                "message": "Excluded adverse wording from source",
                "reason": "Accepted limitation",
            }
        ],
        "omissions": [],
    }
    documents = [
        {
            "id": document_id,
            "filename": "adverse-wording.pdf",
            "sha256": "a" * 64,
            "included_in_exhibits": False,
            "excluded_reason": "Replaced after broker review",
            "excluded_by": reviewer_id,
            "processing_status": "completed",
            "affected_requirements": ["coverage.endorsement"],
        }
    ]
    evidence = [
        {
            "id": source_id,
            "document_id": document_id,
            "text": "Old value",
            "locator": {"kind": "pdf", "page": 1, "line": 2},
        }
    ]
    result, _ = render(snapshot, digest(canonical_json(snapshot)), documents, [], evidence)
    pdf_text = "".join(
        page.extract_text() for page in PdfReader(io.BytesIO(result["pdf"][0])).pages
    )
    assert "adverse-wording.pdf" in pdf_text and "Excluded from exhibits" in pdf_text
    assert "Fictional client call" in pdf_text and "Correction history" in pdf_text
    workbook = openpyxl.load_workbook(io.BytesIO(result["xlsx"][0]))
    assert workbook["Exposure schedule"]["B2"].value == "123456789.1234"
    assert "broker_attestation" in [c.value for c in workbook["Reviewed facts"][2]]
    assert "Replaced after broker review" in [c.value for c in workbook["Document index"][2]]
    workbook.close()
    with zipfile.ZipFile(io.BytesIO(result["zip"][0])) as archive:
        stored = json.loads(archive.read("reviewed-snapshot.json"))
        assert stored["facts"][0]["evidence_type"] == "broker_attestation"


def test_unrecognized_material_content_remains_visible_after_partial_extraction() -> None:
    from app.documents import processing_inventory

    source = blocks(
        "Legal insured name: Partial LLC",
        "A manuscript endorsement excludes all work at the supplied location",
    )
    facts = local_extract(source, "policy")
    assert len(facts) == 1
    inventory = processing_inventory(source, facts)
    assert any(
        "manuscript endorsement" in issue["message"] and issue["source_id"] == str(source[1].id)
        for issue in inventory["issues"]
    )


def test_adverse_prose_containing_application_is_not_mistaken_for_a_heading() -> None:
    from app.documents import processing_inventory

    source = blocks(
        "Legal insured name: Partial LLC", "This application excludes all subcontracted work"
    )
    facts = local_extract(source, "application")
    assert any(
        "excludes all subcontracted work" in issue["message"]
        for issue in processing_inventory(source, facts)["issues"]
    )


def test_truncated_material_value_requires_manual_support() -> None:
    source = blocks("Business operations: Incomplete source wording")
    source[0].locator["truncated"] = True
    fact = local_extract(source, "application")[0]
    assert fact.value is None and fact.availability == "unreadable"
    assert not supports_fact(
        source, "insured.operations", "Incomplete source wording", {}, [source[0].id]
    )


def test_unknown_sections_break_policy_context_and_unmarked_roles_abstain() -> None:
    facts = local_extract(
        blocks(
            "Policy number: OLD",
            "Currency: USD",
            "Each occurrence limit: 1000000",
            "Another policy section whose relationship is unknown",
            "Each occurrence limit: 2000000",
        ),
        "mixed",
    )
    assert facts[-1].availability == "unreadable" and "policy_reference" not in facts[-1].context
    unmarked = local_extract(blocks("Currency: CAD", "Each occurrence limit: 2000000"), "unknown")[
        0
    ]
    assert unmarked.availability == "unreadable" and unmarked.context["role"] == "unspecified"


def test_compatible_monetary_unit_spellings_share_one_canonical_context() -> None:
    context = {
        "basis": "payroll",
        "currency": "CAD",
        "unit": "CAD",
        "period_start": "2027-01-01",
        "period_end": "2027-12-31",
    }
    _, normalized = validate_fact("exposure.amount", "1000", context)
    assert normalized["unit"] == "money" and normalized["currency"] == "CAD"
    assert group_key("exposure.amount", context) == group_key(
        "exposure.amount", {**context, "unit": "money"}
    )
