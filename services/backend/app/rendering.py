"""Original submission templates. Rendering consumes only an immutable reviewed snapshot."""

import io
import json
import zipfile
from html import escape
from typing import Any

import openpyxl
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app.common import canonical_json, digest
from app.domain.cgl import FIELD_LABELS

TEMPLATE_VERSION = "original-cgl-2"


def text(value: Any) -> str:
    return canonical_json(value) if isinstance(value, (dict, list)) else str(value)


def safe_cell(value: Any) -> str:
    rendered = text(value)
    return "'" + rendered if rendered.startswith(("=", "+", "-", "@", "\t", "\r")) else rendered


def context_text(context: dict[str, Any]) -> str:
    return "; ".join(
        key.replace("_", " ").capitalize() + ": " + str(value)
        for key, value in context.items()
        if value is not None and value != "unspecified"
    )


def render(
    snapshot: dict[str, Any],
    revision_hash: str,
    documents: list[dict[str, Any]],
    supporting: list[tuple[str, bytes]],
    evidence: list[dict[str, Any]] | None = None,
) -> tuple[dict[str, tuple[bytes, str, str]], dict[str, Any]]:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor("#18332e")
    story: list[Any] = []

    def paragraph(value: str, style: str = "BodyText") -> None:
        story.append(Paragraph(escape(value).replace("\n", "<br/>"), styles[style]))
        story.append(Spacer(1, 7))

    paragraph("Commercial General Liability", "Title")
    paragraph("Broker submission package · Synthetic review build", "Heading2")
    paragraph(f"Proposed period: {snapshot['intended_start']} to {snapshot['intended_end']}")
    paragraph(
        "This original template summarizes broker-reviewed facts. Policy wording, endorsements, and the attached documents control. This software does not recommend coverage or bind insurance."
    )
    facts = snapshot["facts"]
    sources = {source["id"]: source for source in (evidence or [])}
    document_names = {d["id"]: d["filename"] for d in documents}

    def citation(source_id: str) -> str:
        source = sources.get(source_id)
        if not source:
            return source_id
        locator = source["locator"]
        location = (
            f"page {locator.get('page')}, line {locator.get('line')}"
            if locator["kind"] == "pdf"
            else f"{locator.get('sheet')} {locator.get('range')}"
        )
        return f"{document_names.get(source['document_id'], source['document_id'])} ({location}; source {source_id})"

    for prefix, heading in [
        ("insured.", "Insured overview"),
        ("location.", "Locations"),
        ("operations.", "Operations and classifications"),
        ("exposure.", "GL exposures"),
        ("policy.", "Existing policy"),
        ("coverage.", "Existing and requested coverage"),
        ("loss.", "Loss history"),
    ]:
        paragraph(heading, "Heading2")
        selected = [fact for fact in facts if fact["field"].startswith(prefix)]
        if not selected:
            paragraph(
                "Not supplied or not verified. No coverage or no-loss conclusion is inferred."
            )
        if selected and prefix in {"location.", "exposure.", "coverage.", "loss."}:
            rows = [["Field", "Reviewed value", "Applicability"]]
            for fact in selected:
                value = fact["value"]
                rendered_value = context_text(value) if isinstance(value, dict) else text(value)
                rows.append(
                    [FIELD_LABELS[fact["field"]], rendered_value, context_text(fact["context"])]
                )
            cells = [
                [Paragraph(escape(cell).replace("\n", "<br/>"), styles["BodyText"]) for cell in row]
                for row in rows
            ]
            table = LongTable(cells, colWidths=[115, 180, 227], repeatRows=1, splitInRow=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8efed")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend([table, Spacer(1, 10)])
        for fact in selected:
            if prefix in {"location.", "exposure.", "coverage.", "loss."}:
                paragraph("Evidence for " + FIELD_LABELS[fact["field"]], "Heading3")
            else:
                paragraph(FIELD_LABELS[fact["field"]] + ": " + text(fact["value"]), "Heading3")
            paragraph("Context: " + context_text(fact["context"]))
            paragraph(
                "Evidence: "
                + ("; ".join(citation(ref) for ref in fact["source_ids"]) or "Broker attestation")
                + "; review decision "
                + fact["decision_id"]
            )
            paragraph("Resolution: " + fact["reason"])
            if fact.get("evidence_type") == "broker_attestation":
                paragraph(
                    "Attested by "
                    + str(fact.get("reviewer_id", ""))
                    + "; origin: "
                    + str(fact.get("attestation_origin", ""))
                )
            if fact.get("original_source_ids"):
                paragraph(
                    "Correction history (original evidence): "
                    + "; ".join(citation(ref) for ref in fact["original_source_ids"])
                )
    paragraph("Underwriting narrative", "Heading2")
    names = [text(f["value"]) for f in facts if f["field"] == "insured.legal_name"]
    operations = [text(f["value"]) for f in facts if f["field"] == "insured.operations"]
    paragraph(
        "Insured: "
        + "; ".join(names)
        + ". Operations as verified by the broker: "
        + "; ".join(operations)
        + ". Refer to the exposure, coverage, and loss sections for their specific periods and contexts."
    )
    paragraph("Exceptions and omissions", "Heading2")
    for item in snapshot["exceptions"]:
        paragraph(item["message"] + " Reason: " + str(item.get("reason", "")))
    for item in snapshot["omissions"]:
        paragraph("Omitted review group " + item["group_key"] + ": " + item["reason"])
    for item in snapshot.get("rejections", []):
        paragraph("Rejected unsupported extraction " + item["group_key"] + ": " + item["reason"])
    if not snapshot["exceptions"] and not snapshot["omissions"]:
        paragraph("No accepted exceptions or omitted extracted fields in this revision.")
    paragraph("Supporting document index", "Heading2")
    for document in documents:
        paragraph(f"{document['id']} · {document['filename']} · SHA-256 {document['sha256']}")
        if document.get("excluded_reason"):
            paragraph(
                "Excluded from exhibits: "
                + text(
                    {
                        key: document.get(key)
                        for key in (
                            "excluded_reason",
                            "excluded_by",
                            "excluded_at",
                            "processing_status",
                            "affected_requirements",
                            "replaced_by",
                            "supersedes_id",
                        )
                    }
                )
            )
            for assertion in document.get("source_assertions", []):
                paragraph(
                    "Excluded source assertion: "
                    + assertion["field"]
                    + ": "
                    + text(assertion["value"])
                    + "; "
                    + context_text(assertion["context"])
                )
    paragraph("Reviewed revision SHA-256: " + revision_hash)
    paragraph(
        "Template: "
        + TEMPLATE_VERSION
        + ". Approval is recorded separately against this exact package manifest. Candidate previews are unapproved until broker approval."
    )
    SimpleDocTemplate(
        buffer,
        pagesize=(8.5 * inch, 11 * inch),
        leftMargin=45,
        rightMargin=45,
        topMargin=40,
        bottomMargin=40,
        title="CGL broker submission",
    ).build(story)
    pdf = buffer.getvalue()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Reviewed facts"
    sheet.append(
        [
            "Field",
            "Value (exact text)",
            "Context",
            "Evidence IDs",
            "Decision ID",
            "Reason",
            "Evidence type",
            "Attestation origin",
            "Reviewer",
            "Original evidence (correction history)",
        ]
    )
    for fact in facts:
        sheet.append(
            [
                safe_cell(fact["field"]),
                safe_cell(fact["value"]),
                safe_cell(fact["context"]),
                safe_cell(
                    "; ".join(citation(ref) for ref in fact["source_ids"]) or "Broker attestation"
                ),
                safe_cell(fact["decision_id"]),
                safe_cell(fact["reason"]),
                safe_cell(fact.get("evidence_type", "document")),
                safe_cell(fact.get("attestation_origin") or ""),
                safe_cell(fact.get("reviewer_id", "")),
                safe_cell("; ".join(citation(ref) for ref in fact.get("original_source_ids", []))),
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 32, "B": 55, "C": 65, "D": 55, "E": 40, "F": 60}.items():
        sheet.column_dimensions[column].width = width
    for name, prefix, context_fields in [
        (
            "Exposure schedule",
            "exposure.",
            [
                "basis",
                "unit",
                "currency",
                "period_start",
                "period_end",
                "role",
                "scope",
                "location_id",
            ],
        ),
        ("Coverage schedule", "coverage.", ["role", "policy_reference", "currency", "scope"]),
        ("Location schedule", "location.", ["location_id", "entity_id", "scope"]),
        ("Loss schedule", "loss.", ["currency", "period_start", "period_end", "as_of", "scope"]),
    ]:
        schedule = workbook.create_sheet(name)
        claim_fields = (
            [
                "claim_id",
                "carrier",
                "policy",
                "loss_date",
                "valuation",
                "paid",
                "reserves",
                "incurred",
            ]
            if prefix == "loss."
            else []
        )
        schedule.append(
            [
                "Field",
                "Reviewed value",
                *context_fields,
                *claim_fields,
                "Evidence type",
                "Supporting sources",
                "Attestation origin",
                "Decision ID",
            ]
        )
        for fact in facts:
            if not fact["field"].startswith(prefix):
                continue
            claim = fact["value"] if isinstance(fact["value"], dict) else {}
            schedule.append(
                [
                    safe_cell(value)
                    for value in [
                        FIELD_LABELS[fact["field"]],
                        context_text(claim) if claim else fact["value"],
                        *[fact["context"].get(key) or "" for key in context_fields],
                        *[
                            claim.get(key) if claim.get(key) is not None else ""
                            for key in claim_fields
                        ],
                        fact.get("evidence_type", "document"),
                        "; ".join(citation(ref) for ref in fact["source_ids"]),
                        fact.get("attestation_origin") or "",
                        fact["decision_id"],
                    ]
                ]
            )
        schedule.freeze_panes = "A2"
        schedule.auto_filter.ref = schedule.dimensions
        for column in schedule.columns:
            schedule.column_dimensions[column[0].column_letter].width = 26
    notes = workbook.create_sheet("Exceptions")
    notes.append(["Requirement", "Message", "Reason"])
    for item in snapshot["exceptions"]:
        message = str(item["message"])
        for start in range(0, max(1, len(message)), 30000):
            notes.append(
                [
                    safe_cell(item["key"]),
                    safe_cell(message[start : start + 30000]),
                    safe_cell(item.get("reason", "")),
                ]
            )
    for item in snapshot["omissions"]:
        notes.append([safe_cell(item["group_key"]), "Omitted", safe_cell(item["reason"])])
    for item in snapshot.get("rejections", []):
        notes.append(
            [
                safe_cell(item["group_key"]),
                "Rejected unsupported extraction",
                safe_cell(item["reason"]),
            ]
        )
    index = workbook.create_sheet("Document index")
    index.append(
        [
            "Document ID",
            "Filename",
            "SHA-256",
            "Included in exhibits",
            "Exclusion reason",
            "Excluded by",
            "Excluded at",
            "Affected requirements",
            "Replaced by",
            "Supersedes",
            "Processing status",
        ]
    )
    for doc in documents:
        index.append(
            [
                safe_cell(doc["id"]),
                safe_cell(doc["filename"]),
                doc["sha256"],
                str(doc.get("included_in_exhibits", True)),
                *[
                    safe_cell(doc.get(key) or "")
                    for key in (
                        "excluded_reason",
                        "excluded_by",
                        "excluded_at",
                        "affected_requirements",
                        "replaced_by",
                        "supersedes_id",
                        "processing_status",
                    )
                ],
            ]
        )
    source_sheet = workbook.create_sheet("Source references")
    source_sheet.append(["Source ID", "Document", "Location", "Quoted text"])
    for source in evidence or []:
        source_sheet.append(
            [
                source["id"],
                safe_cell(document_names.get(source["document_id"], source["document_id"])),
                safe_cell(citation(source["id"])),
                safe_cell(source["text"]),
            ]
        )
    metadata_sheet = workbook.create_sheet("Source metadata")
    metadata_sheet.append(["Source ID", "Part (join in order)", "Locator JSON"])
    for source in evidence or []:
        locator_json = canonical_json(source["locator"])
        for start in range(0, len(locator_json), 30000):
            metadata_sheet.append(
                [source["id"], start // 30000 + 1, safe_cell(locator_json[start : start + 30000])]
            )
    excluded_sheet = workbook.create_sheet("Excluded assertions")
    excluded_sheet.append(
        ["Document", "Field", "Source value", "Context", "Availability", "Source IDs"]
    )
    for document in documents:
        if document.get("excluded_reason"):
            for assertion in document.get("source_assertions", []):
                excluded_sheet.append(
                    [
                        safe_cell(value)
                        for value in [
                            document["filename"],
                            assertion["field"],
                            assertion["value"],
                            assertion["context"],
                            assertion["availability"],
                            assertion["source_ids"],
                        ]
                    ]
                )
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    workbook.close()
    xlsx = xlsx_buffer.getvalue()
    manifest = {
        "template_version": TEMPLATE_VERSION,
        "revision_hash": revision_hash,
        "documents": documents,
        "source_index": evidence or [],
        "files": [
            {"filename": name, "sha256": digest(data)}
            for name, data in [("submission.pdf", pdf), ("submission.xlsx", xlsx), *supporting]
        ],
    }
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("submission.pdf", pdf)
        archive.writestr("submission.xlsx", xlsx)
        archive.writestr("manifest.json", json.dumps(manifest, indent=2))
        archive.writestr("reviewed-snapshot.json", canonical_json(snapshot))
        for name, data in supporting:
            archive.writestr(name, data)
    result = {
        "pdf": (pdf, "submission.pdf", "application/pdf"),
        "xlsx": (
            xlsx,
            "submission.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        "zip": (zip_buffer.getvalue(), "submission.zip", "application/zip"),
    }
    # The outer approval manifest also binds the archive, without a self-referential archive hash.
    manifest["artifacts"] = [
        {"format": format, "sha256": digest(data), "byte_size": len(data)}
        for format, (data, _, _) in result.items()
    ]
    return result, manifest
